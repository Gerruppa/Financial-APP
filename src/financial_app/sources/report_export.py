"""Writing a report to an XLSX file (issue #240, parent #239), so the Sell Summary and Wyniki leave the app as data.

``report_to_xlsx`` turns a ``domain/report.py`` ``Report`` into a workbook, a sheet per section: amounts and
quantities are numbers, percents are fractions shown as percent, dates are real dates shown day-first. The report
functions build the model for a year and a View Scope from the database and name the file after the report, the year
and, for one Account, the Account.
"""

import io
from collections.abc import Mapping, Sequence
from datetime import date
from decimal import Decimal

import openpyxl
from openpyxl.cell import Cell as SheetCell
from openpyxl.worksheet.worksheet import Worksheet
from sqlalchemy import Engine

from financial_app.domain.portfolio_history import DayValue
from financial_app.domain.report import Cell, Percent, Report, Section, results_report, sale_summary_report
from financial_app.persistence.accounts import list_accounts
from financial_app.persistence.bond_data import load_cpi
from financial_app.persistence.instruments import list_instruments
from financial_app.persistence.transactions import transactions_in_entry_order
from financial_app.sources.export import FORMULA_START, XLSX_DATE_FORMAT, safe_file_part

XLSX_NUMBER_FORMAT = "#,##0.00##"
XLSX_PERCENT_FORMAT = "0.00%"
# A sheet's name holds at most 31 characters and none of these
_SHEET_NAME_LENGTH = 31
_SHEET_NAME_UNSAFE = str.maketrans("", "", "[]:*?/\\")


def report_to_xlsx(report: Report) -> bytes:
    workbook = openpyxl.Workbook()
    first = workbook.active
    assert first is not None
    workbook.properties.title = report.title
    if not report.sections:
        first.title = _sheet_name(report.title)
        for note in report.notes:
            first.append([note])
    for index, section in enumerate(report.sections):
        sheet = first if index == 0 else workbook.create_sheet()
        _write_section(sheet, section)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _write_section(sheet: Worksheet, section: Section) -> None:
    sheet.title = _sheet_name(section.title)
    sheet.append(list(section.headers))
    for row_number, row in enumerate(section.rows, start=2):
        for column, source in enumerate(row, start=1):
            _format(sheet.cell(row_number, column, _value(source)), source)
    if section.notes:
        sheet.append([])
        for note in section.notes:
            sheet.append([note])


def _value(cell: Cell) -> str | int | Decimal | date | None:
    return cell.fraction if isinstance(cell, Percent) else cell


def _format(cell: SheetCell, source: Cell) -> None:
    if isinstance(source, str) and source.startswith(FORMULA_START):
        cell.data_type = "s"  # stays text, not a formula
    elif isinstance(source, Percent):
        cell.number_format = XLSX_PERCENT_FORMAT
    elif isinstance(source, date):
        cell.number_format = XLSX_DATE_FORMAT
    elif isinstance(source, Decimal):
        cell.number_format = XLSX_NUMBER_FORMAT


def _sheet_name(title: str) -> str:
    return title.translate(_SHEET_NAME_UNSAFE)[:_SHEET_NAME_LENGTH]


def export_sale_summary(engine: Engine, year: int, account_id: int | None) -> tuple[str, bytes]:
    """The file name and bytes of the Sell Summary of ``year`` for the View Scope (``account_id``, ``None``: Total)."""
    accounts = {account.id: account.name for account in list_accounts(engine)}
    instruments = {instrument.id: instrument.name for instrument in list_instruments(engine)}
    report = sale_summary_report(transactions_in_entry_order(engine), year, account_id, instruments, accounts)
    return _named(report, year, account_id, accounts), report_to_xlsx(report)


def export_results(
    engine: Engine, year: int, account_id: int | None, history: Sequence[DayValue], today: date
) -> tuple[str, bytes]:
    """The file name and bytes of Wyniki over ``year`` for the View Scope, from ``history``'s days."""
    accounts = list_accounts(engine)
    report = results_report(history, accounts, account_id, year, today, load_cpi(engine))
    return _named(report, year, account_id, {account.id: account.name for account in accounts}), report_to_xlsx(report)


def _named(report: Report, year: int, account_id: int | None, accounts: Mapping[int, str]) -> str:
    name = f"{report.name}-{year}"
    if account_id is not None:
        name += f"-{safe_file_part(accounts[account_id])}"
    return f"{name}.xlsx"
