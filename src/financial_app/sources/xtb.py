"""Importing XTB statements and reconciling them with the Transactions entered (spec 7, L-17).

Nothing is saved until the user has seen the Reconciliation and applied it: ``preview_xtb`` only reads, and
``apply_xtb`` saves every choice or none. Each XTB account is matched to an Account (or left out), each ticker to an
Instrument; both are remembered in the app settings for the next statement. Resolved broker Transactions are not
offered again, so the same files can be imported once more after the next export.
"""

import warnings
import zipfile
from collections.abc import Mapping, Sequence
from dataclasses import dataclass, field, replace
from datetime import date
from pathlib import Path

import openpyxl
from sqlalchemy import Engine

from financial_app.domain.accounts import Account, AccountType
from financial_app.domain.currencies import MissingNbpRateError, NbpRate
from financial_app.domain.formatting import format_date
from financial_app.domain.instruments import Instrument
from financial_app.domain.quantity_corrections import Fix, Shortfall, quantity_shortfalls
from financial_app.domain.reconciliation import Choice, Reconciliation, item_key, reconcile
from financial_app.domain.transactions import PLN, TransactionDraft, TransactionError
from financial_app.domain.xtb_import import (
    SHEET_NAME,
    EntryContext,
    EntryKind,
    XtbEntry,
    XtbError,
    XtbStatement,
    match_instrument,
    parse_statement,
    unresolved,
    xtb_draft,
    xtb_entries,
)
from financial_app.persistence.accounts import list_accounts
from financial_app.persistence.instruments import list_instruments
from financial_app.persistence.reconciliation import (
    Action,
    AddBroker,
    DeleteMine,
    KeepMine,
    SkipBroker,
    TakeBroker,
    after_actions,
    apply_reconciliation,
    handled_broker_ids,
    incoming_transfers,
    unreconciled_transactions,
    with_corrections,
)
from financial_app.persistence.settings import XTB_ACCOUNT, XTB_TICKER, load_setting, save_setting
from financial_app.sources.spreadsheet import NbpLookup

# The remembered match of an XTB account the user left out
LEFT_OUT = "skip"
# The Account Type an XTB product is, to suggest the one Account of that type
PRODUCT_TYPES = {"IKE": AccountType.IKE, "IKZE": AccountType.IKZE}


@dataclass(frozen=True)
class XtbDecisions:
    """The user's choices: an Account per XTB account number (None: not imported), an Instrument per ticker, and how
    each Reconciliation item is resolved (``Choice.BROKER`` unless chosen, the broker being the source of truth)."""

    accounts: Mapping[str, int | None] = field(default_factory=dict)
    instruments: Mapping[str, int] = field(default_factory=dict)
    choices: Mapping[str, Choice] = field(default_factory=dict)
    # The fix of each Sell left uncovered (``XtbPreview.shortfalls``), by ``shortfall_key``
    corrections: Mapping[str, Fix] = field(default_factory=dict)


@dataclass(frozen=True)
class Problem:
    """An entry that cannot become a Transaction, e.g. without its NBP Rate; it blocks applying."""

    entry: XtbEntry
    message: str


@dataclass(frozen=True)
class XtbPreview:
    statements: list[XtbStatement]
    # Each XTB account number's Account, chosen, remembered or suggested; None when it is left out
    accounts: dict[str, int | None]
    unmatched_accounts: list[str]
    # Tickers with no Instrument, with the instrument name XTB gives
    unmatched_tickers: dict[str, str]
    problems: list[Problem]
    unsupported: list[XtbEntry]
    handled: int
    reconciliation: Reconciliation
    choices: Mapping[str, Choice]
    # The Sells and Security Transfers the choices leave uncovered, each with its fix (#278)
    shortfalls: list[Shortfall] = field(default_factory=list)

    @property
    def blocked(self) -> bool:
        unresolved = any(shortfall.fix is None for shortfall in self.shortfalls)
        return bool(self.unmatched_accounts or self.unmatched_tickers or self.problems or unresolved)

    def choice(self, key: str) -> Choice:
        return self.choices.get(key, Choice.BROKER)

    def actions(self) -> list[Action]:
        """What applying saves: each item as chosen, with the quantity corrections."""
        return with_corrections(self.chosen_actions(), self.shortfalls)

    def chosen_actions(self) -> list[Action]:
        """Each Reconciliation item as chosen, before the quantity corrections."""
        actions: list[Action] = []
        for pair in self.reconciliation.pairs:
            if pair.linked:
                # The user's Cash Transfer is XTB's Deposit or Withdrawal and also moves the other Account's cash
                actions.append(KeepMine(pair.mine.id, pair.broker.external_id))
            elif self.choice(pair.key) is Choice.BROKER:
                draft = replace(pair.broker, comment=pair.mine.comment or pair.broker.comment)
                actions.append(TakeBroker(pair.mine.id, draft))
            else:
                actions.append(KeepMine(pair.mine.id, pair.broker.external_id))
        for draft in self.reconciliation.broker_only:
            key = item_key(draft)
            actions.append(AddBroker(draft) if self.choice(key) is Choice.BROKER else SkipBroker(key))
        for transaction in self.reconciliation.mine_only:
            if self.choice(item_key(transaction)) is Choice.BROKER:
                actions.append(DeleteMine(transaction.id))
            else:
                actions.append(KeepMine(transaction.id))
        return actions


def read_statements(path: Path) -> list[XtbStatement]:
    """The statements in an XLSX file, or in every XLSX file of a folder; raises XtbError if one cannot be read."""
    files = sorted(path.glob("*.xlsx")) if path.is_dir() else [path]
    if not files:
        raise XtbError(f"W folderze {path.name} nie ma plików .xlsx.")
    statements = [_read(file) for file in files]
    numbers = [statement.account for statement in statements]
    if len(set(numbers)) != len(numbers):
        raise XtbError("Dwa pliki dotyczą tego samego rachunku XTB; wskaż jeden.")
    return statements


def _read(path: Path) -> XtbStatement:
    try:
        # Not read-only: xStation's files declare a one-cell size, which the read-only mode believes. They also have no
        # default style, which openpyxl warns about and replaces with its own
        with warnings.catch_warnings():
            warnings.filterwarnings("ignore", "Workbook contains no default style", UserWarning)
            workbook = openpyxl.load_workbook(path, data_only=True)
    except OSError, ValueError, zipfile.BadZipFile:
        raise XtbError(f"Nie udało się otworzyć pliku {path.name}. Wskaż zestawienie XTB (.xlsx).") from None
    try:
        if SHEET_NAME not in workbook.sheetnames:
            raise XtbError(f"W pliku {path.name} nie ma zakładki „{SHEET_NAME}”.")
        currency = _account_currency(workbook)
        if currency not in ("", PLN):
            raise XtbError(f"Rachunek XTB w pliku {path.name} jest w {currency}; import obsługuje rachunki w PLN.")
        return parse_statement(workbook[SHEET_NAME].iter_rows(values_only=True))
    finally:
        workbook.close()


def preview_xtb(
    engine: Engine, statements: Sequence[XtbStatement], rates: NbpLookup, decisions: XtbDecisions
) -> XtbPreview:
    """Match the statements' accounts and tickers and reconcile their entries with the Transactions entered, without
    saving anything."""
    accounts = {account.id: account for account in list_accounts(engine)}
    instruments = {instrument.id: instrument for instrument in list_instruments(engine)}
    chosen = {s.account: _account_for(engine, s, accounts, decisions) for s in statements}
    unmatched_accounts = [number for number, choice in chosen.items() if isinstance(choice, _Unmatched)]
    decided = {number: choice for number, choice in chosen.items() if not isinstance(choice, _Unmatched)}
    mapped = {number: choice for number, choice in decided.items() if choice is not None}
    included = [statement for statement in statements if statement.account in mapped]

    handled = handled_broker_ids(engine)
    every = xtb_entries(included)
    entries = [left for entry in every if (left := unresolved(entry, handled)) is not None]
    done = len(every) - len(entries)
    unsupported = [entry for entry in entries if entry.kind is EntryKind.UNSUPPORTED]

    tickers = {entry.ticker: entry.name for entry in entries if entry.needs_instrument}
    resolved = {ticker: _instrument_for(engine, ticker, instruments, decisions) for ticker in tickers}
    unmatched_tickers = {ticker: tickers[ticker] for ticker, found in resolved.items() if found is None}

    drafts: list[TransactionDraft] = []
    problems = []
    nbp_cache: dict[tuple[str, date], NbpRate | MissingNbpRateError] = {}
    for entry in entries:
        if entry.kind is EntryKind.UNSUPPORTED or (entry.needs_instrument and resolved[entry.ticker] is None):
            continue
        account = accounts[mapped[entry.account]]
        instrument = resolved.get(entry.ticker) if entry.needs_instrument else None
        try:
            nbp_rate = _nbp(rates, instrument, entry.day, nbp_cache)
            target = mapped.get(entry.target) if entry.target else None
            context = EntryContext(account.id, target, instrument, nbp_rate, account.fx_conversion_fee_percent)
            drafts.append(xtb_draft(entry, context))
        except TransactionError as error:
            problems.append(Problem(entry, str(error)))

    ranges = {mapped[s.account]: (s.first_day, s.last_day) for s in included}
    reconciliation = reconcile(drafts, unreconciled_transactions(engine, ranges), incoming_transfers(engine, ranges))
    preview = XtbPreview(
        list(statements),
        decided,
        unmatched_accounts,
        unmatched_tickers,
        problems,
        unsupported,
        done,
        reconciliation,
        dict(decisions.choices),
    )
    if unmatched_accounts or unmatched_tickers:
        return preview
    known = after_actions(engine, preview.chosen_actions())
    return replace(preview, shortfalls=quantity_shortfalls(known, decisions.corrections, trim_saved=False))


def apply_xtb(
    engine: Engine, statements: Sequence[XtbStatement], rates: NbpLookup, decisions: XtbDecisions, today: date
) -> XtbPreview:
    """Save every choice of the Reconciliation, all or none, and remember the matches; returns what was applied.

    Raises XtbError if anything is still unmatched or an entry cannot become a Transaction; nothing is saved then.
    """
    preview = preview_xtb(engine, statements, rates, decisions)
    if preview.unmatched_accounts or preview.unmatched_tickers:
        raise XtbError("Wybierz konto dla każdego rachunku XTB i instrument dla każdego tickera.")
    if preview.problems:
        first = preview.problems[0]
        raise XtbError(f"Operacja XTB z {format_date(first.entry.day)}: {first.message}")
    if any(shortfall.fix is None for shortfall in preview.shortfalls):
        raise XtbError("Wybierz poprawkę dla każdej sprzedaży bez pokrycia.")
    apply_reconciliation(engine, preview.actions(), today)
    for number, account_id in preview.accounts.items():
        save_setting(engine, XTB_ACCOUNT + number, LEFT_OUT if account_id is None else str(account_id))
    for ticker, instrument_id in decisions.instruments.items():
        save_setting(engine, XTB_TICKER + ticker, str(instrument_id))
    return preview


def _account_currency(workbook: openpyxl.Workbook) -> str:
    """The account's currency as the Open Positions sheet's summary gives it ("" for an IKE or IKZE, which leave it
    empty); the Cash Operations do not say."""
    if "Open Positions" not in workbook.sheetnames:
        return ""
    for row in workbook["Open Positions"].iter_rows(values_only=True):
        if len(row) > 3 and row[1] == "Open position value" and row[3]:
            return str(row[3])
    return ""


class _Unmatched:
    pass


_UNMATCHED = _Unmatched()


def _account_for(
    engine: Engine, statement: XtbStatement, accounts: Mapping[int, Account], decisions: XtbDecisions
) -> int | None | _Unmatched:
    """The Account an XTB account goes into: the user's choice, else the remembered one, else the one active Account
    of its product's type (IKE, IKZE); None when it is left out."""
    if statement.account in decisions.accounts:
        return decisions.accounts[statement.account]
    saved = load_setting(engine, XTB_ACCOUNT + statement.account)
    if saved == LEFT_OUT:
        return None
    if saved is not None and int(saved) in accounts:
        return int(saved)
    account_type = PRODUCT_TYPES.get(statement.product)
    of_type = [a.id for a in accounts.values() if a.active and a.account_type is account_type]
    return of_type[0] if account_type is not None and len(of_type) == 1 else _UNMATCHED


def _instrument_for(
    engine: Engine, ticker: str, instruments: Mapping[int, Instrument], decisions: XtbDecisions
) -> Instrument | None:
    chosen = decisions.instruments.get(ticker)
    if chosen is None and (saved := load_setting(engine, XTB_TICKER + ticker)) is not None:
        chosen = int(saved)
    if chosen is not None and chosen in instruments:
        return instruments[chosen]
    return match_instrument(ticker, instruments.values())


def _nbp(
    rates: NbpLookup,
    instrument: Instrument | None,
    day: date,
    cache: dict[tuple[str, date], NbpRate | MissingNbpRateError],
) -> NbpRate | None:
    if instrument is None or instrument.quote_currency == PLN:
        return None
    key = (instrument.quote_currency, day)
    if key not in cache:
        try:
            cache[key] = rates.before(*key)
        except MissingNbpRateError as error:
            cache[key] = error
    found = cache[key]
    if isinstance(found, MissingNbpRateError):
        raise XtbError(str(found)) from None
    return found
