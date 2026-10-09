"""Importing Transactions from the user's copy of the Inwestomat sheet (spec 7, issue #20).

Nothing is saved until the user has seen the preview and approved it: ``preview_import`` only reads, and
``save_import`` saves all the rows or none. Rows already imported (same fingerprint) are skipped, so the same file can
be imported again.
"""

import zipfile
from collections.abc import Mapping, Sequence
from dataclasses import dataclass, field
from datetime import date
from pathlib import Path
from typing import Literal, Protocol

import openpyxl
from sqlalchemy import Engine

from financial_app.domain.currencies import MissingNbpRateError, NbpRate
from financial_app.domain.formatting import format_date
from financial_app.domain.instruments import Instrument, InstrumentDraft
from financial_app.domain.spreadsheet_import import (
    SheetError,
    SheetLine,
    TickerKind,
    classify_ticker,
    fingerprints,
    parse_lines,
    sheet_draft,
    similar,
)
from financial_app.domain.transactions import PLN, Origin, TransactionDraft, TransactionError
from financial_app.persistence.accounts import list_accounts
from financial_app.persistence.instruments import (
    add_asset_class,
    add_instrument,
    delete_instrument,
    list_asset_classes,
    list_instruments,
)
from financial_app.persistence.transactions import add_transactions, imported_external_ids, list_transactions

SHEET_NAME = "Transakcje"
# The Instrument of a draft in a preview whose Instrument does not exist yet; never saved
_UNSAVED_INSTRUMENT = 0


class NbpLookup(Protocol):
    """Where the NBP Rate comes from: ``sources.nbp.NbpRates`` or a test double."""

    def before(self, currency: str, day: date) -> NbpRate: ...


@dataclass(frozen=True)
class Decisions:
    """What the user chose for what the sheet does not match: an Account per sheet Konto, and for each Instrument
    name its existing Instrument's id, or None to create it."""

    accounts: Mapping[str, int] = field(default_factory=dict)
    instruments: Mapping[str, int | None] = field(default_factory=dict)
    # Fingerprints of rows the user chose to import although a Transaction entered by hand looks the same
    keep_similar: frozenset[str] = frozenset()


@dataclass(frozen=True)
class PreviewLine:
    """A sheet row as it would be saved: ``ok``, skipped as already ``duplicate``, held back as ``similar`` to a
    Transaction entered by hand (unless the user keeps it), or ``error`` with the reason."""

    line: SheetLine
    status: Literal["ok", "duplicate", "similar", "error"]
    message: str
    draft: TransactionDraft | None
    external_id: str = ""


@dataclass(frozen=True)
class Preview:
    lines: list[PreviewLine]
    # Sheet Konto names with no Account chosen, and Instrument names with no existing Instrument matched or chosen
    unmatched_accounts: list[str]
    unmatched_instruments: list[str]
    # Instruments that saving creates: the first sheet row of each name, for its Asset Class and currency
    new_instruments: dict[str, SheetLine]

    @property
    def blocked(self) -> bool:
        return any(preview_line.status == "error" for preview_line in self.lines)

    @property
    def to_save(self) -> int:
        return sum(preview_line.status == "ok" for preview_line in self.lines)

    @property
    def duplicates(self) -> int:
        return sum(preview_line.status == "duplicate" for preview_line in self.lines)

    @property
    def held_back(self) -> list[PreviewLine]:
        return [preview_line for preview_line in self.lines if preview_line.status == "similar"]


def read_sheet_lines(path: Path) -> list[SheetLine]:
    """The rows of the Transakcje tab of the workbook at ``path``; raises SheetError if it cannot be read."""
    try:
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except OSError, ValueError, zipfile.BadZipFile:
        raise SheetError(f"Nie udało się otworzyć arkusza {path.name}. Wskaż plik .xlsx.") from None
    try:
        if SHEET_NAME not in workbook.sheetnames:
            raise SheetError(f"W pliku nie ma zakładki „{SHEET_NAME}”.")
        rows = workbook[SHEET_NAME].iter_rows(values_only=True)
        header = next(rows, ())
        return parse_lines(header, rows)
    finally:
        workbook.close()


def preview_import(engine: Engine, lines: Sequence[SheetLine], rates: NbpLookup, decisions: Decisions) -> Preview:
    """Check every row against the saved Accounts, Instruments and imports, without saving anything."""
    accounts = {account.name.casefold(): account.id for account in list_accounts(engine)}
    instruments = {instrument.name.casefold(): instrument for instrument in list_instruments(engine)}
    saved = imported_external_ids(engine)
    by_hand = [t for t in list_transactions(engine) if t.origin is Origin.MANUAL]
    prints = fingerprints(lines)
    nbp_cache: dict[tuple[str, date], NbpRate | MissingNbpRateError] = {}

    unmatched_accounts = sorted({line.account for line in lines if _account_id(line, accounts, decisions) is None})
    names = [classify_ticker(line.ticker).name for line in lines if classify_ticker(line.ticker).needs_instrument]
    unmatched_instruments = sorted({name for name in names if _instrument_id(name, instruments, decisions) is None})
    new_instruments: dict[str, SheetLine] = {}
    for line in lines:
        ticker = classify_ticker(line.ticker)
        if ticker.needs_instrument and _instrument_id(ticker.name, instruments, decisions) is None:
            new_instruments.setdefault(ticker.name, line)

    preview_lines = []
    for line, external_id in zip(lines, prints, strict=True):
        if external_id in saved:
            preview_lines.append(PreviewLine(line, "duplicate", "Już zaimportowana.", None, external_id))
            continue
        try:
            draft = _draft_for(line, external_id, accounts, instruments, decisions, rates, nbp_cache)
        except TransactionError as error:
            preview_lines.append(PreviewLine(line, "error", str(error), None, external_id))
            continue
        match = next((t for t in by_hand if similar(draft, t)), None)
        if match is not None and external_id not in decisions.keep_similar:
            message = f"Podobna do ręcznej transakcji z {format_date(match.date)} (nr {match.id})."
            preview_lines.append(PreviewLine(line, "similar", message, draft, external_id))
        else:
            preview_lines.append(PreviewLine(line, "ok", "", draft, external_id))
    return Preview(preview_lines, unmatched_accounts, unmatched_instruments, new_instruments)


def save_import(engine: Engine, lines: Sequence[SheetLine], rates: NbpLookup, decisions: Decisions) -> Preview:
    """Create the new Instruments and save every ``ok`` row, all or none; returns the preview that was saved.

    Raises SheetError, naming the first row to fix, if any row is blocked; nothing is saved then.
    """
    preview = preview_import(engine, lines, rates, decisions)
    if preview.blocked:
        first = next(line for line in preview.lines if line.status == "error")
        raise SheetError(f"Wiersz {first.line.row}: {first.message} Popraw arkusz albo wybór i spróbuj ponownie.")
    created: list[int] = []
    try:
        class_ids = {asset_class.name.casefold(): asset_class.id for asset_class in list_asset_classes(engine)}
        for name, line in preview.new_instruments.items():
            class_key = line.asset_class.casefold()
            if class_key not in class_ids:
                class_ids[class_key] = add_asset_class(engine, line.asset_class).id
            is_bond = classify_ticker(line.ticker).kind is TickerKind.BOND
            draft = InstrumentDraft(
                name,
                class_ids[class_key],
                line.currency,
                manual_price=line.price if is_bond else None,
            )
            created.append(add_instrument(engine, draft).id)
        # Now the new Instruments exist, so the drafts can name them
        final = preview_import(engine, lines, rates, decisions)
        add_transactions(engine, [item.draft for item in final.lines if item.status == "ok" and item.draft is not None])
    except Exception:
        # Any failure, a database one included, leaves no new Instrument behind
        for instrument_id in created:
            delete_instrument(engine, instrument_id)
        raise
    return final


def _account_id(line: SheetLine, accounts: Mapping[str, int], decisions: Decisions) -> int | None:
    if line.account in decisions.accounts:
        return decisions.accounts[line.account]
    return accounts.get(line.account.casefold())


def _instrument_id(name: str, instruments: Mapping[str, Instrument], decisions: Decisions) -> int | None:
    """The existing Instrument the name resolves to, or None: then it is new unless ``decisions`` names one."""
    if (instrument := instruments.get(name.casefold())) is not None:
        return instrument.id
    return decisions.instruments.get(name)


def _draft_for(
    line: SheetLine,
    external_id: str,
    accounts: Mapping[str, int],
    instruments: Mapping[str, Instrument],
    decisions: Decisions,
    rates: NbpLookup,
    nbp_cache: dict[tuple[str, date], NbpRate | MissingNbpRateError],
) -> TransactionDraft:
    account_id = _account_id(line, accounts, decisions)
    if account_id is None:
        raise SheetError(f"Wybierz konto dla „{line.account}”.")
    ticker = classify_ticker(line.ticker)
    instrument_id = None
    nbp_rate = None
    if ticker.needs_instrument:
        instrument_id = _instrument_id(ticker.name, instruments, decisions)
        if instrument_id is None:
            if not line.asset_class:
                raise SheetError(f"Brak klasy aktywów dla nowego instrumentu „{ticker.name}”.")
            instrument_id = _UNSAVED_INSTRUMENT
        existing = instruments.get(ticker.name.casefold())
        if existing is not None and existing.quote_currency != line.currency:
            raise SheetError(
                f"Instrument „{ticker.name}” ma walutę {existing.quote_currency}, a arkusz {line.currency}."
            )
        if line.currency != PLN:
            nbp_rate = _nbp(rates, line, nbp_cache)
    draft = sheet_draft(
        line,
        account_id=account_id,
        instrument_id=instrument_id,
        nbp_rate=nbp_rate,
        external_id=external_id,
    )
    return draft


def _nbp(rates: NbpLookup, line: SheetLine, cache: dict[tuple[str, date], NbpRate | MissingNbpRateError]) -> NbpRate:
    key = (line.currency, line.day)
    if key not in cache:
        try:
            cache[key] = rates.before(line.currency, line.day)
        except MissingNbpRateError as error:
            cache[key] = error
    found = cache[key]
    if isinstance(found, MissingNbpRateError):
        raise SheetError(str(found)) from None
    return found
