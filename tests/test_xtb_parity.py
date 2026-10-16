"""The user's XTB exports (L-17) read whole: each account's units per ticker must equal XTB's own Open Positions
sheet and its cash the Cash Operations total. Skipped without ``reference/`` (git-ignored, personal data)."""

from collections import defaultdict
from decimal import Decimal
from pathlib import Path

import openpyxl
import pytest

from financial_app.domain.xtb_import import EntryKind, XtbEntry, xtb_entries
from financial_app.sources.xtb import read_statements

EXPORTS = sorted(path for path in (Path(__file__).parent.parent / "reference").glob("*/*.xlsx") if "_" in path.name)

pytestmark = [
    pytest.mark.skipif(not EXPORTS, reason="no XTB exports in reference/"),
    pytest.mark.filterwarnings("ignore:Workbook contains no default style"),
]


def open_positions(path: Path) -> dict[str, Decimal]:
    """The units per ticker of the Open Positions sheet's summary rows (those with a Category)."""
    book = openpyxl.load_workbook(path, data_only=True)
    units = {}
    header = False
    for row in book["Open Positions"].iter_rows(values_only=True):
        if row[1] == "Instrument/Position":
            header = True
        elif header and row[3]:
            units[str(row[2])] = Decimal(str(row[5]))
    return units


def cash_total(path: Path) -> Decimal:
    book = openpyxl.load_workbook(path, data_only=True)
    total = next(row[5] for row in book["Cash Operations"].iter_rows(values_only=True) if row[0] == "Total")
    return Decimal(str(total)).quantize(Decimal("0.01"))


@pytest.mark.parametrize("path", EXPORTS, ids=lambda path: path.name.split("_")[0])
def test_units_and_cash_agree_with_xtb(path: Path) -> None:
    [statement] = read_statements(path)
    entries = xtb_entries([statement])
    assert [entry.note for entry in entries if entry.kind is EntryKind.UNSUPPORTED] == []
    assert held(entries) == open_positions(path)
    assert cash(entries) == cash_total(path)


def held(entries: list[XtbEntry]) -> dict[str, Decimal]:
    units: dict[str, Decimal] = defaultdict(Decimal)
    for entry in entries:
        if entry.kind is EntryKind.BUY:
            units[entry.ticker] += entry.quantity or 0
        elif entry.kind is EntryKind.SELL:
            units[entry.ticker] -= entry.quantity or 0
    return {ticker: quantity for ticker, quantity in units.items() if quantity}


def cash(entries: list[XtbEntry]) -> Decimal:
    signs = {
        EntryKind.DEPOSIT: 1,
        EntryKind.SELL: 1,
        EntryKind.INTEREST: 1,
        EntryKind.DIVIDEND: 1,
        EntryKind.WITHDRAWAL: -1,
        EntryKind.BUY: -1,
        EntryKind.TRANSFER: -1,
    }
    total = Decimal(0)
    for entry in entries:
        total += signs[entry.kind] * entry.amount
        if entry.kind in (EntryKind.INTEREST, EntryKind.DIVIDEND):
            total -= entry.withholding_tax
    return total
