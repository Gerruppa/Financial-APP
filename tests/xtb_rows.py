"""Synthetic XTB statements in the layout of xStation's account-history XLSX (L-17); the user's own exports stay in
``reference/``, which is git-ignored."""

from collections.abc import Sequence
from datetime import datetime
from pathlib import Path

import openpyxl

from financial_app.domain.xtb_import import COLUMNS, SHEET_NAME

UNTIL = datetime(2026, 10, 13, 8, 35, 51)
HEADER = ["Type", "Instrument", "Ticker", "Category", "Time", "Amount", "ID", "Comment", "Product", "Position ID"]
assert set(COLUMNS.values()) <= set(HEADER)

Row = tuple[str, str, str, datetime, float, str, str, str]


def op(
    kind: str,
    time: datetime,
    amount: float,
    id: str,
    comment: str = "",
    ticker: str = "",
    instrument: str = "",
    position: str = "",
) -> Row:
    return (kind, instrument, ticker, time, amount, id, comment, position)


def sheet_rows(account: str, operations: Sequence[Row], product: str = "IKE") -> list[tuple[object, ...]]:
    rows: list[tuple[object, ...]] = [
        ("Account number", account),
        ("Cash Operations", ""),
        ("Date from (UTC)", datetime(2006, 1, 1)),
        ("Date to (UTC)", UNTIL),
        tuple(HEADER),
    ]
    for kind, instrument, ticker, time, amount, id, comment, position in operations:
        category = "ETF" if ticker else ""
        rows.append((kind, instrument, ticker, category, time, amount, id, comment, product, position))
    rows.append(("Total", None, None, None, None, 0.0, None, None, None, None))
    return rows


def write_statement(path: Path, account: str, operations: Sequence[Row], product: str = "IKE") -> Path:
    """An XLSX with the three sheets xStation exports, the Cash Operations one filled."""
    book = openpyxl.Workbook()
    closed = book.active
    assert closed is not None
    closed.title = "Closed Positions"
    closed.append(["Account number", account])
    cash = book.create_sheet(SHEET_NAME)
    for row in sheet_rows(account, operations, product):
        cash.append(list(row))
    book.create_sheet("Open Positions").append(["Account number", account])
    book.save(path)
    return path
