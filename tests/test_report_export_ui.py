"""The "Eksportuj XLSX" buttons of the Sell Summary and Wyniki tabs (issue #240): the year choice, the View Scope and
the download; the figures are tested with the report model and the files."""

import io
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

import openpyxl
import pytest
from nicegui import ui
from nicegui.testing import User, user_simulation
from sqlalchemy import Engine

from financial_app.domain.accounts import AccountDraft, AccountType
from financial_app.domain.instruments import InstrumentDraft
from financial_app.domain.transactions import TransactionDraft, TransactionType, buy_or_sell
from financial_app.persistence.accounts import add_account
from financial_app.persistence.db import init_db
from financial_app.persistence.instruments import add_instrument, list_asset_classes
from financial_app.persistence.transactions import add_transactions
from financial_app.sources.history import QuoteHistory
from financial_app.sources.nbp import NbpRates
from financial_app.sources.portfolio_history import PortfolioHistory
from financial_app.ui.shell import build_shell

TODAY = date(2026, 10, 12)
NOW = datetime(2026, 10, 12, 9, 0)


@pytest.fixture
def engine(tmp_path: Path) -> Engine:
    """IKE: PZU bought 10 at 50 on 5.01.2026 and sold at 60 on 4.03.2026; XTB: only a Deposit."""
    engine = init_db(tmp_path / "app.sqlite3")
    ike = add_account(engine, AccountDraft("IKE", "mBank", AccountType.IKE, ("PLN",))).id
    xtb = add_account(engine, AccountDraft("XTB", "XTB", AccountType.REGULAR, ("PLN",))).id
    shares = next(c for c in list_asset_classes(engine) if c.name == "Akcje polskie")
    pzu = add_instrument(engine, InstrumentDraft("PZU", shares.id, "PLN")).id
    add_transactions(
        engine,
        [
            TransactionDraft(ike, date(2026, 1, 2), TransactionType.DEPOSIT, Decimal(1000)),
            TransactionDraft(xtb, date(2026, 1, 2), TransactionType.DEPOSIT, Decimal(1000)),
            buy_or_sell(ike, date(2026, 1, 5), TransactionType.BUY, pzu, Decimal(10), Decimal(50)),
            buy_or_sell(ike, date(2026, 3, 4), TransactionType.SELL, pzu, Decimal(10), Decimal(60)),
        ],
    )
    return engine


async def _inline(function, *args):  # type: ignore[no-untyped-def]
    return function(*args)


def _shell(engine: Engine):  # type: ignore[no-untyped-def]
    quotes = QuoteHistory(engine, lambda url: "Brak danych", lambda: TODAY, lambda: NOW, lambda url: None)
    history = PortfolioHistory(engine, quotes, NbpRates(engine), lambda: TODAY, _inline)
    return lambda: build_shell(engine, now=lambda: NOW, history=history)


def _choose_year(user: User, marker: str, year: int) -> None:
    [select] = user.find(marker=marker).elements
    assert isinstance(select, ui.select)
    select.set_value(year)


async def test_the_sell_summary_exports_the_last_finished_year_by_default(engine: Engine) -> None:
    async with user_simulation(_shell(engine)) as user:
        await user.open("/sprzedaz")
        user.find(marker="sale-export").click()
        response = await user.download.next()

    sheet = openpyxl.load_workbook(io.BytesIO(response.content)).active
    assert sheet is not None
    assert "Brak danych" in sheet["A1"].value  # nothing was sold in 2025


async def test_the_sell_summary_exports_the_chosen_year_and_the_scope(engine: Engine) -> None:
    async with user_simulation(_shell(engine)) as user:
        await user.open("/sprzedaz")
        _choose_year(user, "sale-export-year", 2026)
        user.find(marker="view-scope").click()
        await user.should_see(marker="scope-menu")
        user.find(marker="scope-1").click()
        user.find(marker="sale-export").click()
        response = await user.download.next()

    sheet = openpyxl.load_workbook(io.BytesIO(response.content))["Sprzedaże"]
    assert [cell.value for cell in sheet[2]][:2] == ["PZU", "IKE"]


async def test_wyniki_exports_the_chosen_year(engine: Engine) -> None:
    async with user_simulation(_shell(engine)) as user:
        await user.open("/wyniki")
        await user.should_see(marker="results-twr")
        _choose_year(user, "results-export-year", 2026)
        user.find(marker="results-export").click()
        response = await user.download.next()

    sheet = openpyxl.load_workbook(io.BytesIO(response.content))["Wyniki roku"]
    assert sheet["A3"].value == "TWR"
