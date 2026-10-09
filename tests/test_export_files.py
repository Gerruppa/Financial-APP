"""The export files (spec 5, issue #22): CSV that Polish Excel opens with its characters, and XLSX read back."""

import io
from datetime import date, datetime
from decimal import Decimal

import openpyxl

from financial_app.sources.export import to_csv, to_xlsx

HEADERS = ("Nr", "Data", "Konto", "Komentarz", "Kwota")
ROWS: list[tuple[object, ...]] = [
    (1, date(2026, 3, 7), "IKE", "zakup ąęśćżźńół, zł", Decimal("1234.50")),
    (2, date(2026, 3, 8), "IKZE", None, Decimal("-0.5")),
]


def test_csv_is_utf8_with_bom_semicolons_and_polish_numbers() -> None:
    data = to_csv(HEADERS, ROWS)

    assert data.startswith(b"\xef\xbb\xbf")
    lines = data.decode("utf-8-sig").splitlines()
    assert lines[0] == "Nr;Data;Konto;Komentarz;Kwota"
    assert lines[1] == "1;07.03.2026;IKE;zakup ąęśćżźńół, zł;1234,50"
    assert lines[2] == "2;08.03.2026;IKZE;;-0,5"


def test_csv_quotes_a_field_containing_the_delimiter() -> None:
    data = to_csv(HEADERS, [(1, date(2026, 3, 7), "IKE", "a;b", Decimal(1))])

    assert '1;07.03.2026;IKE;"a;b";1' in data.decode("utf-8-sig")


def test_xlsx_round_trips_polish_text_numbers_and_dates() -> None:
    data = to_xlsx(HEADERS, ROWS)

    sheet = openpyxl.load_workbook(io.BytesIO(data)).active
    assert sheet is not None
    read = [tuple(cell.value for cell in row) for row in sheet.iter_rows()]
    assert read[0] == HEADERS
    assert read[1][3] == "zakup ąęśćżźńół, zł"
    assert read[1][4] == 1234.5
    assert read[1][1] == datetime(2026, 3, 7)
    assert read[2][3] is None
    assert read[2][4] == -0.5


def test_xlsx_shows_dates_day_first() -> None:
    sheet = openpyxl.load_workbook(io.BytesIO(to_xlsx(HEADERS, ROWS))).active
    assert sheet is not None
    assert sheet["B2"].number_format == "DD.MM.YYYY"


def test_csv_neutralises_text_that_a_spreadsheet_would_run_as_a_formula() -> None:
    data = to_csv(
        HEADERS, [(1, date(2026, 3, 7), "IKE", "=1+1", Decimal(1)), (2, date(2026, 3, 7), "IKE", "-5 zł", Decimal(-5))]
    )

    lines = data.decode("utf-8-sig").splitlines()
    assert lines[1].split(";")[3] == "'=1+1"
    assert lines[2].split(";")[3] == "'-5 zł"
    assert lines[2].split(";")[4] == "-5"


def test_xlsx_keeps_text_that_starts_like_a_formula_as_text() -> None:
    sheet = openpyxl.load_workbook(io.BytesIO(to_xlsx(HEADERS, [(1, date(2026, 3, 7), "IKE", "=1+1", None)]))).active
    assert sheet is not None
    assert sheet["D2"].value == "=1+1"
    assert sheet["D2"].data_type == "s"
