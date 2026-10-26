"""The report XLSX files (issue #240): numbers and dates stay Excel cells, the name carries report, year and Account."""

import io
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

import openpyxl
import pytest
from openpyxl.worksheet.worksheet import Worksheet
from sqlalchemy import Engine

from financial_app.domain.accounts import AccountDraft, AccountType
from financial_app.domain.instruments import InstrumentDraft
from financial_app.domain.portfolio_history import DayValue
from financial_app.domain.report import NOTHING_TO_REPORT, Percent, Report, Section
from financial_app.domain.transactions import TransactionDraft, TransactionType, buy_or_sell
from financial_app.persistence.accounts import add_account
from financial_app.persistence.db import init_db
from financial_app.persistence.instruments import add_instrument, list_asset_classes
from financial_app.persistence.transactions import add_transaction
from financial_app.sources.report_export import export_results, export_sale_summary, report_to_xlsx


@pytest.fixture
def engine(tmp_path: Path) -> Engine:
    return init_db(tmp_path / "app.sqlite3")


def _sheet(data: bytes, title: str) -> Worksheet:
    return openpyxl.load_workbook(io.BytesIO(data))[title]


def _sell_once(engine: Engine, account_name: str = "mBank IKE") -> int:
    account = add_account(engine, AccountDraft(account_name, "mBank", AccountType.IKE, ("PLN",))).id
    shares = next(c for c in list_asset_classes(engine) if c.name == "Akcje polskie")
    instrument = add_instrument(engine, InstrumentDraft("PZU", shares.id, "PLN")).id
    add_transaction(engine, TransactionDraft(account, date(2025, 1, 2), TransactionType.DEPOSIT, Decimal(1000)))
    for kind, day, price in (
        (TransactionType.BUY, date(2025, 1, 5), "50"),
        (TransactionType.SELL, date(2025, 3, 6), "60"),
    ):
        add_transaction(engine, buy_or_sell(account, day, kind, instrument, Decimal(10), Decimal(price), Decimal(0)))
    return account


def test_the_sell_summary_file_holds_numbers_and_dates_not_text(engine: Engine) -> None:
    _sell_once(engine)

    name, data = export_sale_summary(engine, 2025, None)

    assert name == "podsumowanie-sprzedazy-2025.xlsx"
    sheet = _sheet(data, "Sprzedaże")
    assert [cell.value for cell in sheet[1]][:4] == ["Instrument", "Konto", "Zakup", "Sprzedaż"]
    row = [cell.value for cell in sheet[2]]
    assert row[:2] == ["PZU", "mBank IKE"]
    assert row[2] == datetime(2025, 1, 5)
    assert sheet["C2"].number_format == "DD.MM.YYYY"
    assert row[5:8] == [500, 600, 100]  # cost, revenue, result
    assert sheet["I2"].number_format == "0.00%"
    assert row[8] == pytest.approx(0.2)


def test_one_account_names_the_file_with_its_safe_name(engine: Engine) -> None:
    account = _sell_once(engine, 'Konto: "A/B"')

    name, _ = export_sale_summary(engine, 2025, account)

    assert name == "podsumowanie-sprzedazy-2025-Konto_ _A_B_.xlsx"


def test_an_empty_year_gives_a_valid_file_saying_there_is_nothing_to_report(engine: Engine) -> None:
    _sell_once(engine)

    _, data = export_sale_summary(engine, 2023, None)

    sheet = openpyxl.load_workbook(io.BytesIO(data)).active
    assert sheet is not None
    assert sheet["A1"].value == NOTHING_TO_REPORT


def test_the_results_file_follows_the_year_and_scope(engine: Engine) -> None:
    account = _sell_once(engine)
    days = [
        DayValue(date(2025, 1, 2), {account: Decimal(1000)}, {}, Decimal(1000), {account: Decimal(1000)}, frozenset()),
        DayValue(
            date(2025, 12, 31), {account: Decimal(1100)}, {}, Decimal(1000), {account: Decimal(1000)}, frozenset()
        ),
    ]

    name, data = export_results(engine, 2025, account, days, date(2026, 10, 26))

    assert name == "wyniki-2025-mBank IKE.xlsx"
    sheet = _sheet(data, "Wyniki roku")
    values = {row[0].value: row[1].value for row in sheet.iter_rows(min_row=2)}
    assert values["TWR"] == pytest.approx(0.1)
    assert sheet["B3"].number_format == "0.00%"


def test_text_starting_with_a_formula_character_stays_text() -> None:
    report = Report("r", "R", (Section("S", ("A", "B"), (("=SUM(1)", Percent(Decimal("0.5"))),)),))

    sheet = _sheet(report_to_xlsx(report), "S")

    assert sheet["A2"].data_type == "s"
    assert sheet["A2"].value == "=SUM(1)"
